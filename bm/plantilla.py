"""Plantilla Excel de registro operativo generada con el catálogo actual (desplegables siempre al día).
Mismas hojas y columnas que la plantilla que ya usa el hotel."""

from __future__ import annotations

import io
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MARINO = PatternFill("solid", fgColor="14284B")
FILAS = 500

INSTRUCCIONES = [
    ("Registro operativo de desayuno · BM", True),
    ("", False),
    ("1. Hoja Registro: una fila por plato. Misma Fecha = un único desayuno en BM.", False),
    ("2. Huespedes: 1 = comensal nuevo, 0 o vacío = mismo comensal pidiendo otro plato.", False),
    ("3. Nombre: elija del desplegable. Cantidad = raciones de ese plato.", False),
    ("4. Extra1..4: añadidos (Bacon, Aguacate...). Huevo pochado/cocido/revuelto SUSTITUYE al frito;", False),
    ("   Tostada integral / sin gluten SUSTITUYE al pan de la ficha (mismas rebanadas).", False),
    ("5. Omitir1/2: quitar algo de la ficha (Sin huevo, Sin tostada o un ingrediente).", False),
    ("6. Bebidas de desayuno, comida, cena y buffet tienen su propia hoja.", False),
    ("7. En BM: Registrar > Importar Excel. Primero se ve una vista previa; nada se guarda hasta confirmar.", False),
    ("", False),
    ("Volver a importar un día lo SUSTITUYE (no duplica). No borre filas de un día a medias antes de reimportarlo.", True),
]


def _hoja(wb, nombre, columnas, anchos):
    ws = wb.create_sheet(nombre)
    ws.append(columnas)
    for i, ancho in enumerate(anchos, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = MARINO, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    for r in range(2, FILAS + 1):
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
    return ws


def _lista(ws, rango_celdas, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    for r in rango_celdas.split():
        dv.add(r)


def generar(con: sqlite3.Connection) -> bytes:
    q = lambda sql: [r[0] for r in con.execute(sql)]  # noqa: E731
    listas = {
        "Recetas_desayuno": q("SELECT nombre FROM recetas WHERE activo=1 AND servicio='desayuno' ORDER BY nombre"),
        "Extras": q("SELECT etiqueta FROM atajos WHERE activo=1 AND grupo IN ('extra','leche') ORDER BY etiqueta"),
        "Omitir": q("SELECT etiqueta FROM atajos WHERE activo=1 AND grupo IN ('omitir','extra') ORDER BY grupo DESC, etiqueta"),
        "Bebidas_desayuno": q("""SELECT nombre FROM recetas WHERE activo=1 AND servicio='bebidas'
                                 UNION SELECT etiqueta FROM atajos WHERE activo=1 AND grupo IN ('bebida','leche') ORDER BY 1"""),
        "Recetas_comida": q("SELECT nombre FROM recetas WHERE activo=1 AND servicio='comida' ORDER BY nombre"),
        "Recetas_cena": q("SELECT nombre FROM recetas WHERE activo=1 AND servicio IN ('cena','comida') ORDER BY nombre"),
        "Buffet": q("SELECT etiqueta FROM atajos WHERE activo=1 AND grupo='buffet' ORDER BY etiqueta"),
    }
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    ins.column_dimensions["A"].width = 110
    for texto, negrita in INSTRUCCIONES:
        ins.append([texto])
        ins.cell(row=ins.max_row, column=1).font = Font(bold=negrita, size=13 if ins.max_row == 1 else 11)

    cat = wb.create_sheet("Catalogo")
    ref = {}
    for i, (nombre, valores) in enumerate(listas.items(), start=1):
        col = get_column_letter(i)
        cat.cell(row=1, column=i, value=nombre).font = Font(bold=True)
        cat.column_dimensions[col].width = 32
        for j, v in enumerate(valores, start=2):
            cat.cell(row=j, column=i, value=v)
        ref[nombre] = f"Catalogo!${col}$2:${col}${max(len(valores) + 1, 2)}"

    reg = _hoja(wb, "Registro", ["Fecha", "Huespedes", "Tipo", "Nombre", "Cantidad", "Extra1", "Cant1", "Extra2", "Cant2",
                                 "Extra3", "Cant3", "Extra4", "Cant4", "Omitir1", "Omitir2", "Notas"],
                [12, 10, 10, 30, 9, 18, 6, 18, 6, 18, 6, 18, 6, 18, 18, 30])
    _lista(reg, f"B2:B{FILAS}", '"0,1"')
    _lista(reg, f"C2:C{FILAS}", '"Receta,Extra,Producto"')
    _lista(reg, f"D2:D{FILAS}", ref["Recetas_desayuno"])
    _lista(reg, " ".join(f"{c}2:{c}{FILAS}" for c in "FHJL"), ref["Extras"])
    _lista(reg, f"N2:O{FILAS}", ref["Omitir"])
    for nombre, lista in (("RegistroBebidasDesayuno", "Bebidas_desayuno"), ("RegistroComida", "Recetas_comida"),
                          ("RegistroCena", "Recetas_cena")):
        ws = _hoja(wb, nombre, ["Fecha", "Tipo", "Nombre", "Cantidad", "Notas"], [12, 10, 32, 9, 30])
        _lista(ws, f"B2:B{FILAS}", '"Receta,Extra,Producto"')
        _lista(ws, f"C2:C{FILAS}", ref[lista])
    buf = _hoja(wb, "ConsumoBuffet", ["Fecha", "Concepto", "Cantidad", "Motivo", "Notas"], [12, 34, 9, 14, 30])
    _lista(buf, f"B2:B{FILAS}", ref["Buffet"])
    _lista(buf, f"D2:D{FILAS}", '"Consumo,Merma,Expiración,Limpieza"')
    wb.move_sheet("Catalogo", offset=len(wb.sheetnames))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
