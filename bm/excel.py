"""Excel operativo de desayuno (misma plantilla que ya usa el hotel) -> consumos.

Hojas: Registro (platos de desayuno), RegistroBebidasDesayuno, RegistroComida, RegistroCena, ConsumoBuffet.
- Cada (hoja, día) es UN consumo. Reimportar un día lo sustituye (nunca duplica).
- Un día con algún error no se importa; el resto sí. Siempre se puede previsualizar antes.
- Extras/omisiones se resuelven con la tabla `atajos`; huevo y pan se sustituyen por grupo.
"""

from __future__ import annotations

import re
import sqlite3
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

from bm import costing

HOJAS = {  # hoja -> (servicio, grupos de atajos válidos como línea suelta)
    "Registro": ("desayuno", ("extra", "leche", "bebida")),
    "RegistroBebidasDesayuno": ("desayuno", ("bebida", "leche", "extra")),
    "RegistroComida": ("comida", ("extra",)),
    "RegistroCena": ("cena", ("extra",)),
    "ConsumoBuffet": ("desayuno", ("buffet",)),
}


def _n(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s.lower()).split())


def _num(v, defecto=None):
    if v in (None, ""):
        return defecto
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return defecto


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


@dataclass
class Fila:
    hoja: str
    fila: int
    fecha: date
    tipo: str
    nombre: str
    cantidad: float
    huesped: int = 0
    extras: list = field(default_factory=list)   # [(etiqueta, cant)]
    omitir: list = field(default_factory=list)   # [etiqueta]
    motivo: str = "consumo"
    nota: str = ""


def leer(path) -> tuple[list[Fila], list[tuple[str, int, str]]]:
    """-> (filas válidas, errores de lectura [(hoja, fila, mensaje)])"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    filas, errores = [], []
    for hoja in HOJAS:
        if hoja not in wb.sheetnames:
            continue
        it = wb[hoja].iter_rows(values_only=True)
        cab = {_n(c): i for i, c in enumerate(next(it, ()) or ()) if c}

        def col(row, nombre, _cab=cab):
            i = _cab.get(nombre)
            return row[i] if i is not None and i < len(row) else None

        for n_fila, row in enumerate(it, start=2):
            nombre = str(col(row, "nombre") or col(row, "concepto") or "").strip()
            if not nombre:
                continue
            f = _fecha(col(row, "fecha"))
            if not f:
                errores.append((hoja, n_fila, f"Fecha no válida: {col(row, 'fecha')!r}"))
                continue
            tipo = str(col(row, "tipo") or "").strip()
            hues = _num(col(row, "huespedes"), None)
            if tipo in ("0", "1", "0.0", "1.0"):  # costumbre del hotel: 0/1 en Tipo = huésped
                hues = hues if hues is not None else float(tipo)
                tipo = "Receta"
            extras = []
            for i in range(1, 5):
                e = str(col(row, f"extra{i}") or "").strip()
                if e:
                    extras.append((e, _num(col(row, f"cant{i}"), 1) or 1))
            motivo = _n(col(row, "motivo") or "consumo")
            filas.append(Fila(
                hoja=hoja, fila=n_fila, fecha=f, tipo=_n(tipo) or ("buffet" if hoja == "ConsumoBuffet" else "receta"),
                nombre=nombre, cantidad=_num(col(row, "cantidad"), 1) or 0, huesped=1 if hues and hues >= 1 else 0,
                extras=extras, omitir=[str(x).strip() for x in (col(row, "omitir1"), col(row, "omitir2")) if x],
                motivo="consumo" if motivo in ("", "consumo") else motivo, nota=str(col(row, "notas") or "").strip(),
            ))
    return filas, errores


class Resolver:
    """Traduce filas del Excel a líneas de producto con los datos actuales (recetas, atajos, sustituciones)."""

    def __init__(self, con: sqlite3.Connection):
        self.recetas = {_n(r["nombre"]): r["id"] for r in con.execute("SELECT id, nombre FROM recetas WHERE activo=1")}
        self.porciones = dict(con.execute("SELECT id, porciones FROM recetas").fetchall())
        self.ingredientes = defaultdict(list)
        for r in con.execute("SELECT receta_id, producto, cantidad FROM receta_lineas"):
            self.ingredientes[r["receta_id"]].append((r["producto"], r["cantidad"]))
        self.atajos = defaultdict(dict)
        for a in con.execute("SELECT * FROM atajos WHERE activo=1"):
            self.atajos[a["grupo"]][_n(a["etiqueta"])] = dict(a)
        self.grupo_de = {r["producto"]: r["grupo"] for r in con.execute("SELECT * FROM sustitucion")}
        self.del_dia = {(_n(r["etiqueta"]), r["dia_semana"]): r["receta_id"] for r in con.execute("SELECT * FROM recetas_dia")}
        self.productos = {}
        for p in con.execute("SELECT codigo, nombre FROM productos WHERE activo=1 AND es_tpv=0 AND (categoria LIKE '1%' OR origen='bm2')"):
            self.productos[_n(p["nombre"])] = p["codigo"]
            self.productos[_n(p["codigo"])] = p["codigo"]
        self.nombre_producto = {v: k for k, v in self.productos.items()}

    def receta(self, nombre: str, fecha: date) -> str | None:
        k = _n(nombre)
        return self.del_dia.get((k, fecha.weekday())) or self.recetas.get(k)

    def lineas_receta(self, rid: str, raciones: float) -> list[list]:
        f = raciones / (self.porciones.get(rid) or 1)
        return [[p, q * f, rid] for p, q in self.ingredientes[rid]]

    def atajo(self, etiqueta: str, grupos) -> dict | None:
        k = _n(etiqueta)
        return next((self.atajos[g][k] for g in grupos if k in self.atajos[g]), None)

    def producto(self, nombre: str) -> str | None:
        k = _n(nombre)
        if k in self.productos:
            return self.productos[k]
        tokens = k.split()
        hits = {c for n, c in self.productos.items() if len(tokens) >= 2 and all(t in n.split() for t in tokens)}
        return hits.pop() if len(hits) == 1 else None

    def _cantidad_atajo(self, a: dict, veces: float) -> list[list]:
        if a["receta_id"]:
            return self.lineas_receta(a["receta_id"], a["cantidad"] * veces)
        return [[a["producto"], a["cantidad"] * veces, None]] if a["producto"] else []

    def suelto(self, nombre: str, cantidad: float, grupos) -> list[list]:
        a = self.atajo(nombre, grupos)
        if a:
            return self._cantidad_atajo(a, cantidad)
        p = self.producto(nombre)
        if p:
            return [[p, cantidad, None]]
        raise ValueError(f"No reconozco «{nombre}» (ni atajo ni producto)")

    def fila(self, f: Fila) -> list[list]:
        servicio, grupos = HOJAS[f.hoja]
        if f.cantidad <= 0:
            raise ValueError("Cantidad 0 o vacía")
        if f.hoja == "ConsumoBuffet":
            return self.suelto(f.nombre, f.cantidad, ("buffet",))
        rid = self.receta(f.nombre, f.fecha) if f.tipo in ("receta", "") else None
        if not rid:
            if f.tipo == "receta" and not (self.atajo(f.nombre, grupos) or self.producto(f.nombre)):
                raise ValueError(f"Receta no encontrada: «{f.nombre}»")
            return self.suelto(f.nombre, f.cantidad, grupos)

        lineas = self.lineas_receta(rid, f.cantidad)
        extras_grupos = grupos + ("omitir",)
        for etq in f.omitir:  # primero quitar de la ficha...
            a = self.atajo(etq, extras_grupos)
            if a and a["sustituye"]:
                lineas = [l for l in lineas if self.grupo_de.get(l[0]) != a["sustituye"]]
                continue
            p = (a or {}).get("producto") or self.producto(etq) or next(
                (l[0] for l in lineas if _n(etq) in self.nombre_producto.get(l[0], "")), None)
            if not p or not any(l[0] == p for l in lineas):
                raise ValueError(f"Omitir «{etq}»: no está en la receta")
            lineas = [l for l in lineas if l[0] != p]
        for etq, k in f.extras:  # ...luego añadir extras o sustituir huevo/pan
            a = self.atajo(etq, extras_grupos)
            if a and a["sustituye"] and any(self.grupo_de.get(l[0]) == a["sustituye"] for l in lineas):
                quitadas = [l for l in lineas if self.grupo_de.get(l[0]) == a["sustituye"]]
                lineas = [l for l in lineas if l not in quitadas]
                piezas = sum(q / self._unidad(p, a["sustituye"]) for p, q, _ in quitadas) or f.cantidad
                lineas += [[x[0], x[1], rid] for x in self._cantidad_atajo(a, piezas * k)]
            elif a and a["sustituye"] and not a["producto"]:
                continue  # "Sin huevo" puesto como extra: no hay nada que quitar
            else:
                lineas += self.suelto(etq, k, extras_grupos)
        return lineas

    def _unidad(self, producto: str, grupo: str) -> float:
        """Tamaño de 1 pieza (1 huevo, 1 rebanada) según el atajo del grupo para ese producto."""
        for a in self.atajos["extra"].values():
            if a["sustituye"] == grupo and a["producto"] == producto and a["cantidad"]:
                return a["cantidad"]
        return 1.0


# Registros heredados de BM v2 que este Excel sustituye para el mismo día (vienen del mismo Excel).
REEMPLAZA_BM2 = {
    "Registro": ("desayuno-xlsx", "import-ago26-desayuno", "desayuno-media", "desayuno-1208"),
    "RegistroBebidasDesayuno": ("bebidas-desayuno",),
    "ConsumoBuffet": ("import-ago26-buffet", "buffet-xlsx"),
}


def _ref(hoja: str, fecha: date, tipo: str) -> str:
    return f"excel:{hoja}:{fecha.isoformat()}:{tipo}"


def _previos(con, hoja: str, fecha: date, tipo: str) -> list[int]:
    ids = [r[0] for r in con.execute("SELECT id FROM consumos WHERE ref=?", (_ref(hoja, fecha, tipo),))]
    if hoja == "ConsumoBuffet" and tipo == "consumo":  # el buffet confirmado en BM ese día también se sustituye
        ids += [r[0] for r in con.execute("SELECT id FROM consumos WHERE ref=?", (f"buffet:{fecha.isoformat()}",))]
    if tipo == "consumo":
        for pre in REEMPLAZA_BM2.get(hoja, ()):
            ids += [r[0] for r in con.execute(
                "SELECT id FROM consumos WHERE origen='bm2' AND tipo='consumo' AND fecha=? AND ref LIKE ?",
                (fecha.isoformat(), f"bm2:%:{pre}%"))]
    return ids


def planificar(con: sqlite3.Connection, filas: list[Fila], errores_lectura=()) -> list[dict]:
    """Agrupa por (hoja, día, consumo|merma) y resuelve cada fila. No escribe nada."""
    from bm.consumos import precio_actual

    res = Resolver(con)
    grupos: dict[tuple, dict] = {}
    for f in filas:
        tipo = "consumo" if f.motivo == "consumo" else "merma"
        g = grupos.setdefault((f.hoja, f.fecha, tipo), {
            "hoja": f.hoja, "fecha": f.fecha.isoformat(), "tipo": tipo, "servicio": HOJAS[f.hoja][0],
            "comensales": 0, "filas": 0, "lineas": [], "errores": [], "notas": set()})
        g["filas"] += 1
        g["comensales"] += f.huesped
        if f.nota and not f.nota.upper().startswith("OK "):
            g["notas"].add(f.nota)
        if tipo == "merma":
            g["notas"].add(f"Motivo: {f.motivo}")
        try:
            g["lineas"] += res.fila(f)
        except ValueError as e:
            g["errores"].append({"fila": f.fila, "mensaje": str(e), "nombre": f.nombre})
    plan = []
    cache = {}
    for (hoja, fecha, tipo), g in sorted(grupos.items(), key=lambda x: (x[0][1], x[0][0])):
        coste = sum(q * (cache.setdefault(p, precio_actual(con, p, g["fecha"])) or 0) for p, q, _ in g["lineas"])
        plan.append({**g, "notas": "; ".join(sorted(g["notas"])) or None, "n_lineas": len(g["lineas"]),
                     "coste_estimado": round(coste, 2), "reemplaza": len(_previos(con, hoja, fecha, tipo)),
                     "comensales": g["comensales"] if hoja == "Registro" else None})
    for hoja, fila, msg in errores_lectura:
        plan.append({"hoja": hoja, "fecha": None, "tipo": "consumo", "servicio": HOJAS[hoja][0], "comensales": None,
                     "filas": 1, "lineas": [], "errores": [{"fila": fila, "mensaje": msg, "nombre": ""}],
                     "notas": None, "n_lineas": 0, "coste_estimado": 0, "reemplaza": 0})
    return plan


def importar(con: sqlite3.Connection, path, usuario: str | None = None, confirmar: bool = False) -> dict:
    filas, errores = leer(path)
    plan = planificar(con, filas, errores)
    if not confirmar:
        return {"plan": [{k: v for k, v in g.items() if k != "lineas"} for g in plan], "importados": 0}
    hechos = 0
    for g in plan:
        if g["errores"] or not g["lineas"]:
            continue
        fecha = date.fromisoformat(g["fecha"])
        for cid in _previos(con, g["hoja"], fecha, g["tipo"]):
            con.execute("DELETE FROM consumos WHERE id=?", (cid,))
        cid = con.execute(
            """INSERT INTO consumos(fecha, servicio, tipo, origen, ref, comensales, nota, usuario)
               VALUES(?,?,?,'excel',?,?,?,?)""",
            (g["fecha"], g["servicio"], g["tipo"], _ref(g["hoja"], fecha, g["tipo"]), g["comensales"],
             f"Excel {g['hoja']} ({g['filas']} filas)" + (f" · {g['notas']}" if g["notas"] else ""), usuario),
        ).lastrowid
        con.executemany(
            "INSERT INTO consumo_lineas(consumo_id, producto, cantidad, receta_id) VALUES(?,?,?,?)",
            [(cid, p, q, r) for p, q, r in g["lineas"] if q > 0],
        )
        hechos += 1
    con.commit()
    costing.valorar(con)
    return {"plan": [{k: v for k, v in g.items() if k != "lineas"} for g in plan], "importados": hechos}
