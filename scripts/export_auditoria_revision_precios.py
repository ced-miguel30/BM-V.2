"""Excel de auditoría: productos que necesitan revisión de precio / unidad.

Lee D:\\work\\2-BM-DATOS y escribe
D:\\work\\auditoria_revision_precios_YYYYMMDD.xlsx

Uso:
  py -3 scripts\\export_auditoria_revision_precios.py
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

HOTEL = Path(r"D:\work\2-BM-DATOS\data\datos_hotel.json")
OUT_DIR = Path(r"D:\work")

# Pack / caja en el nombre (Ud ambiguas)
_PACK_RE = re.compile(
    r"(?i)(\d+\s*UD\b|\d+\s*UND\b|\d+\s*UN\b|\bCAJA\b|\bPACK\b|\bB\/\d+|\b\d+\s*X\s*\d+)",
)

ACCIONES = {
    "redondeo_centimo": "Revisar si Ud es pack; valorar no redondear a 0 en fracciones",
    "cero_con_lote": "Revisar línea a coste 0 pese a tener lote con precio",
    "cero_sin_lote": "Dar de alta compra/lote o corregir consumo sin precio",
    "ud_vs_caja": "Clarificar si la unidad de inventario es pieza o caja/pack",
    "desvio_lote": "Revisar precio de lote / revalorizar consumos",
    "b_no_bebida": "Confirmar clasificación producto vs bebida",
    "receta_sin_lote": "Dar de alta compra/lote o quitar de receta",
    "extremo_alto": "Validar precio unitario y unidad",
    "extremo_bajo": "Validar precio unitario y unidad",
}


def _money(x: float) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _es_pack(nombre: str | None) -> bool:
    return bool(_PACK_RE.search(nombre or ""))


def _load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _prods(data: dict) -> dict[str, dict]:
    return {p["id"]: p for p in (data.get("productos") or []) if p.get("id")}


def _lotes_stats(data: dict) -> dict[str, dict[str, float | int]]:
    """producto_id -> {n_lotes, qty, precio_total, coste_ud_medio}."""
    out: dict[str, dict[str, float | int]] = {}
    for l in data.get("lotes") or []:
        if l.get("anulado"):
            continue
        pid = l.get("producto_id")
        if not pid:
            continue
        q = float(l.get("cantidad") or 0)
        pt = float(l.get("precio_total") or 0)
        st = out.setdefault(pid, {"n_lotes": 0, "qty": 0.0, "precio_total": 0.0})
        st["n_lotes"] = int(st["n_lotes"]) + 1
        st["qty"] = float(st["qty"]) + q
        st["precio_total"] = float(st["precio_total"]) + pt
    for st in out.values():
        q = float(st["qty"])
        st["coste_ud_medio"] = (float(st["precio_total"]) / q) if q > 0 else 0.0
    return out


def _consumo_agg(data: dict) -> dict[str, dict[str, float | int]]:
    """Agrega líneas de desayuno/servicio/merma."""
    agg: dict[str, dict[str, float | int]] = {}
    for key in ("desayunos", "registros_servicio", "mermas"):
        for r in data.get(key) or []:
            if r.get("anulado"):
                continue
            for ln in r.get("lineas") or []:
                pid = ln.get("producto_id")
                if not pid:
                    continue
                q = float(ln.get("cantidad") or 0)
                c = float(ln.get("coste") or 0)
                if q == 0 and c == 0:
                    continue
                a = agg.setdefault(
                    pid,
                    {
                        "qty": 0.0,
                        "coste": 0.0,
                        "n": 0,
                        "n_cero": 0,
                        "qty_cero": 0.0,
                        "min_qty": None,
                        "max_qty": None,
                        "fracciones": 0,
                    },
                )
                a["qty"] = float(a["qty"]) + q
                a["coste"] = float(a["coste"]) + c
                a["n"] = int(a["n"]) + 1
                if q > 0 and c <= 0:
                    a["n_cero"] = int(a["n_cero"]) + 1
                    a["qty_cero"] = float(a["qty_cero"]) + q
                if q > 0:
                    mn = a["min_qty"]
                    mx = a["max_qty"]
                    a["min_qty"] = q if mn is None else min(float(mn), q)
                    a["max_qty"] = q if mx is None else max(float(mx), q)
                    if q < 1:
                        a["fracciones"] = int(a["fracciones"]) + 1
    return agg


def _recetas_por_ing(
    data: dict,
) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    out: dict[str, list[str]] = defaultdict(list)
    qty_ge1: dict[str, list[str]] = defaultdict(list)
    for r in data.get("recetas") or []:
        nombre = r.get("nombre") or r.get("id") or "?"
        for ing in r.get("ingredientes") or []:
            pid = ing.get("producto_id")
            if not pid:
                continue
            if nombre not in out[pid]:
                out[pid].append(nombre)
            q = float(ing.get("cantidad") or 0)
            if q >= 1 and nombre not in qty_ge1[pid]:
                qty_ge1[pid].append(nombre)
    return dict(out), dict(qty_ge1)


def _ejemplos(nombres: list[str] | None, n: int = 3) -> str:
    if not nombres:
        return ""
    return "; ".join(nombres[:n])


def _base_row(pid: str, p: dict, cons: dict | None, lote: dict | None) -> dict[str, Any]:
    cons = cons or {}
    lote = lote or {}
    qty = float(cons.get("qty") or 0)
    coste = float(cons.get("coste") or 0)
    cu = (coste / qty) if qty > 0 else None
    lu = float(lote.get("coste_ud_medio") or 0) or None
    if lu == 0:
        lu = None
    return {
        "producto_id": pid,
        "nombre": p.get("nombre") or "",
        "unidad": p.get("unidad") or "",
        "es_bebida": bool(p.get("es_bebida")),
        "qty_consumo": round(qty, 6) if qty else 0,
        "coste_consumo": round(coste, 2) if coste else 0,
        "coste_ud_consumo": round(cu, 6) if cu is not None else "",
        "coste_ud_lote_medio": round(lu, 6) if lu is not None else "",
        "n_lotes": int(lote.get("n_lotes") or 0),
        "n_lineas_cero": int(cons.get("n_cero") or 0),
    }


def build_sheets(data: dict) -> dict[str, list[dict[str, Any]]]:
    prods = _prods(data)
    lotes = _lotes_stats(data)
    cons = _consumo_agg(data)
    recetas_map, pack_receta_ge1 = _recetas_por_ing(data)

    coste_cero: list[dict] = []
    ud_caja: list[dict] = []
    desvio: list[dict] = []
    b_no_beb: list[dict] = []
    receta_sin: list[dict] = []
    extremos: list[dict] = []

    # 01 COSTE CERO
    for pid, a in cons.items():
        if int(a.get("n_cero") or 0) <= 0:
            continue
        p = prods.get(pid) or {"id": pid, "nombre": pid}
        lote = lotes.get(pid)
        lu = float(lote["coste_ud_medio"]) if lote and float(lote.get("qty") or 0) > 0 else None
        qty_cero = float(a.get("qty_cero") or 0)
        if lu is not None and lu > 0 and _money(qty_cero * lu) == 0:
            flag = "redondeo_centimo"
            prio = "alta"
        elif lote and int(lote.get("n_lotes") or 0) > 0:
            flag = "cero_con_lote"
            prio = "alta"
        else:
            flag = "cero_sin_lote"
            prio = "alta"
        row = _base_row(pid, p, a, lote)
        row.update(
            {
                "prioridad": prio,
                "flag": flag,
                "detalle": f"líneas coste 0: {a['n_cero']}; qty en esas líneas: {qty_cero:.6g}",
                "recetas_ejemplo": _ejemplos(recetas_map.get(pid)),
                "accion_sugerida": ACCIONES[flag],
            }
        )
        coste_cero.append(row)

    # 02 UD VS CAJA
    for pid, p in prods.items():
        if not _es_pack(p.get("nombre")):
            continue
        a = cons.get(pid)
        fracciones = int((a or {}).get("fracciones") or 0)
        en_receta_ge1 = bool(pack_receta_ge1.get(pid))
        if fracciones <= 0 and not en_receta_ge1:
            continue
        motivos = []
        if fracciones > 0:
            motivos.append(f"consumo fraccionario ({fracciones} líneas qty<1)")
        if en_receta_ge1:
            motivos.append(f"receta qty≥1: {_ejemplos(pack_receta_ge1.get(pid))}")
        row = _base_row(pid, p, a, lotes.get(pid))
        row.update(
            {
                "prioridad": "alta",
                "flag": "ud_vs_caja",
                "detalle": "; ".join(motivos),
                "recetas_ejemplo": _ejemplos(recetas_map.get(pid)),
                "accion_sugerida": ACCIONES["ud_vs_caja"],
            }
        )
        ud_caja.append(row)

    # 03 DESVIO LOTE
    for pid, a in cons.items():
        qty = float(a.get("qty") or 0)
        coste = float(a.get("coste") or 0)
        if qty <= 0 or coste <= 0:
            continue
        lote = lotes.get(pid)
        if not lote or float(lote.get("qty") or 0) <= 0:
            continue
        lu = float(lote["coste_ud_medio"])
        if lu <= 0:
            continue
        cu = coste / qty
        ratio = cu / lu
        if 0.5 <= ratio <= 2.0:
            continue
        p = prods.get(pid) or {"nombre": pid}
        row = _base_row(pid, p, a, lote)
        row.update(
            {
                "prioridad": "media",
                "flag": "desvio_lote",
                "detalle": f"ratio cons/lote={ratio:.3f}",
                "recetas_ejemplo": _ejemplos(recetas_map.get(pid)),
                "accion_sugerida": ACCIONES["desvio_lote"],
            }
        )
        desvio.append(row)

    # 04 ID B NO BEBIDA
    for pid, p in prods.items():
        if not str(pid).startswith("b"):
            continue
        if p.get("es_bebida"):
            continue
        row = _base_row(pid, p, cons.get(pid), lotes.get(pid))
        row.update(
            {
                "prioridad": "media",
                "flag": "b_no_bebida",
                "detalle": "id b* con es_bebida=False (posible traspaso Noray)",
                "recetas_ejemplo": _ejemplos(recetas_map.get(pid)),
                "accion_sugerida": ACCIONES["b_no_bebida"],
            }
        )
        b_no_beb.append(row)

    # 05 RECETA SIN LOTE
    for pid, nombres in recetas_map.items():
        if lotes.get(pid) and int(lotes[pid].get("n_lotes") or 0) > 0:
            continue
        p = prods.get(pid)
        if p is None:
            nombre = "PRODUCTO INEXISTENTE"
            unidad = ""
            es_bebida = False
        else:
            nombre = p.get("nombre") or ""
            unidad = p.get("unidad") or ""
            es_bebida = bool(p.get("es_bebida"))
        row = {
            "producto_id": pid,
            "nombre": nombre,
            "unidad": unidad,
            "es_bebida": es_bebida,
            "qty_consumo": "",
            "coste_consumo": "",
            "coste_ud_consumo": "",
            "coste_ud_lote_medio": "",
            "n_lotes": 0,
            "n_lineas_cero": "",
            "prioridad": "alta",
            "flag": "receta_sin_lote",
            "detalle": f"usado en {len(nombres)} receta(s) sin lotes activos",
            "recetas_ejemplo": _ejemplos(nombres),
            "accion_sugerida": ACCIONES["receta_sin_lote"],
        }
        # fill consumo if any
        if pid in cons:
            base = _base_row(pid, p or {"nombre": nombre}, cons[pid], None)
            for k in (
                "qty_consumo",
                "coste_consumo",
                "coste_ud_consumo",
                "n_lineas_cero",
            ):
                row[k] = base[k]
        receta_sin.append(row)

    # 06 EXTREMOS
    for pid, a in cons.items():
        qty = float(a.get("qty") or 0)
        coste = float(a.get("coste") or 0)
        if qty <= 0:
            continue
        cu = coste / qty
        if cu >= 30:
            flag = "extremo_alto"
            prio = "media"
        elif 0 < cu < 0.02:
            flag = "extremo_bajo"
            prio = "baja"
        else:
            continue
        p = prods.get(pid) or {"nombre": pid}
        row = _base_row(pid, p, a, lotes.get(pid))
        row.update(
            {
                "prioridad": prio,
                "flag": flag,
                "detalle": f"coste unitario consumo={cu:.6g}",
                "recetas_ejemplo": _ejemplos(recetas_map.get(pid)),
                "accion_sugerida": ACCIONES[flag],
            }
        )
        extremos.append(row)

    def _sort(rows: list[dict]) -> list[dict]:
        prio_ord = {"alta": 0, "media": 1, "baja": 2}
        return sorted(
            rows,
            key=lambda r: (
                prio_ord.get(str(r.get("prioridad")), 9),
                str(r.get("nombre") or ""),
            ),
        )

    return {
        "01_COSTE_CERO": _sort(coste_cero),
        "02_UD_VS_CAJA": _sort(ud_caja),
        "03_DESVIO_LOTE": _sort(desvio),
        "04_ID_B_NO_BEBIDA": _sort(b_no_beb),
        "05_RECETA_SIN_LOTE": _sort(receta_sin),
        "06_EXTREMOS": _sort(extremos),
    }


_COLS = [
    "prioridad",
    "flag",
    "producto_id",
    "nombre",
    "unidad",
    "es_bebida",
    "qty_consumo",
    "coste_consumo",
    "coste_ud_consumo",
    "coste_ud_lote_medio",
    "n_lotes",
    "n_lineas_cero",
    "detalle",
    "recetas_ejemplo",
    "accion_sugerida",
]


def _write_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append(_COLS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c, "") for c in _COLS])
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 10,
        "B": 16,
        "C": 12,
        "D": 42,
        "E": 8,
        "F": 10,
        "G": 12,
        "H": 12,
        "I": 14,
        "J": 16,
        "K": 10,
        "L": 12,
        "M": 48,
        "N": 36,
        "O": 48,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def write_excel(path: Path, hotel: Path, sheets: dict[str, list[dict]]) -> None:
    wb = Workbook()
    # Resumen
    ws = wb.active
    ws.title = "00_RESUMEN"
    ws.append(["campo", "valor"])
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    ws.append(["generado", datetime.now().isoformat(timespec="seconds")])
    ws.append(["fuente_json", str(hotel)])
    ws.append(["descripcion", "Productos a revisar: precio raro y/o Ud vs caja/pack"])
    ws.append([])
    ws.append(["hoja", "filas", "criterio"])
    for cell in ws[6]:
        cell.font = Font(bold=True)
    criterios = {
        "01_COSTE_CERO": "qty>0 y coste=0 (redondeo / con lote / sin lote)",
        "02_UD_VS_CAJA": "nombre pack/caja + consumo fraccionario o receta qty>=1",
        "03_DESVIO_LOTE": "€/ud consumo vs lote fuera de ratio 0.5-2",
        "04_ID_B_NO_BEBIDA": "id b* con es_bebida=False",
        "05_RECETA_SIN_LOTE": "ingrediente de receta sin lotes activos",
        "06_EXTREMOS": "€/ud consumo >=30 o (0, €/ud <0.02)",
    }
    for name, rows in sheets.items():
        ws.append([name, len(rows), criterios.get(name, "")])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 55

    for name, rows in sheets.items():
        _write_sheet(wb, name, rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    if not HOTEL.is_file():
        print(f"No existe {HOTEL}")
        return 1
    print(f"Leyendo {HOTEL} …")
    data = _load(HOTEL)
    sheets = build_sheets(data)
    stamp = datetime.now().strftime("%Y%m%d")
    out = OUT_DIR / f"auditoria_revision_precios_{stamp}.xlsx"
    write_excel(out, HOTEL, sheets)
    print(f"Escrito {out}")
    for name, rows in sheets.items():
        print(f"  {name}: {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
